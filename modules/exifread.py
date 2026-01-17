import os
import zipfile
from io import BytesIO

# -------- Импорты библиотек --------
# PDF
try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

# DOCX
try:
    from docx import Document
except ImportError:
    Document = None

# XLSX
try:
    import openpyxl
except ImportError:
    openpyxl = None

# PPTX
try:
    import pptx
except ImportError:
    pptx = None

# Аудио
try:
    from mutagen.mp3 import MP3
    from mutagen.flac import FLAC
    from mutagen.oggvorbis import OggVorbis
    from mutagen.mp4 import MP4
except ImportError:
    MP3 = FLAC = OggVorbis = MP4 = None

# Изображения
try:
    import exifread
except ImportError:
    exifread = None

# Pillow (PNG, BMP, GIF, WEBP)
try:
    from PIL import Image
except ImportError:
    Image = None

# HEIC
try:
    import pyheif
except ImportError:
    pyheif = None

# Архивы
try:
    import rarfile
except ImportError:
    rarfile = None

try:
    import py7zr
except ImportError:
    py7zr = None

# Hachoir (любые файлы)
try:
    from hachoir.parser import createParser
    from hachoir.metadata import extractMetadata
except ImportError:
    createParser = extractMetadata = None

# FFprobe для видео
import subprocess
import json

# -------- Функции сканирования --------

def scan_pdf(path):
    if PdfReader is None:
        print("[!] PyPDF2 not installed")
        return
    reader = PdfReader(path)
    print(reader.metadata)

def scan_docx(path):
    if Document is None:
        print("[!] python-docx not installed")
        return
    doc = Document(path)
    props = doc.core_properties
    print(f"Author: {props.author}, Created: {props.created}, Modified: {props.modified}")

def scan_xlsx(path):
    if openpyxl is None:
        print("[!] openpyxl not installed")
        return
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"Sheets: {wb.sheetnames}")

def scan_pptx(path):
    if pptx is None:
        print("[!] python-pptx not installed")
        return
    pres = pptx.Presentation(path)
    print(f"Slides: {len(pres.slides)}")

def scan_zip(path):
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            print(f"{info.filename} | {info.date_time} | {info.file_size} bytes")

def scan_rar(path):
    if rarfile is None:
        print("[!] rarfile not installed")
        return
    rf = rarfile.RarFile(path)
    for info in rf.infolist():
        print(f"{info.filename} | {info.date_time} | {info.file_size} bytes")

def scan_7z(path):
    if py7zr is None:
        print("[!] py7zr not installed")
        return
    with py7zr.SevenZipFile(path, mode='r') as archive:
        for name in archive.getnames():
            print(name)

def scan_mp3(path):
    if MP3 is None:
        print("[!] mutagen not installed")
        return
    audio = MP3(path)
    print("Keys:", audio.keys())
    print("Length:", audio.info.length)

def scan_flac(path):
    if FLAC is None:
        print("[!] mutagen not installed")
        return
    audio = FLAC(path)
    print(audio.pprint())

def scan_ogg(path):
    if OggVorbis is None:
        print("[!] mutagen not installed")
        return
    audio = OggVorbis(path)
    print(audio.pprint())

def scan_mp4_audio(path):
    if MP4 is None:
        print("[!] mutagen not installed")
        return
    audio = MP4(path)
    print(audio.pprint())

def scan_image_exif(path):
    if exifread is None:
        print("[!] exifread not installed")
        return
    with open(path, "rb") as f:
        tags = exifread.process_file(f, details=False)
    if tags:
        for tag, value in tags.items():
            print(f"{tag}: {value}")
    else:
        print("[+] No EXIF metadata found")

def scan_image_pillow(path):
    if Image is None:
        print("[!] Pillow not installed")
        return
    img = Image.open(path)
    print(f"Format: {img.format}, Size: {img.size}, Mode: {img.mode}")
    if img.info:
        for k, v in img.info.items():
            print(f"{k}: {v}")

def scan_heic(path):
    """
    Сканирование HEIC через ffprobe (работает на Termux/Android)
    """
    try:
        cmd = [
            "ffprobe",
            "-v", "quiet",
            "-print_format", "json",
            "-show_format",
            "-show_streams",
            path
        ]
        result = subprocess.run(cmd, capture_output=True, text=True)
        info = json.loads(result.stdout)
        # Выводим основные данные
        if "format" in info:
            fmt = info["format"]
            print(f"Filename: {fmt.get('filename')}")
            print(f"Format: {fmt.get('format_name')}")
            print(f"Duration: {fmt.get('duration')} sec")
            print(f"Size: {fmt.get('size')} bytes")
            if "tags" in fmt:
                print("Tags / Metadata:")
                for k, v in fmt["tags"].items():
                    print(f"  {k}: {v}")
        if "streams" in info:
            for i, stream in enumerate(info["streams"]):
                print(f"\nStream #{i}:")
                for k, v in stream.items():
                    print(f"  {k}: {v}")
    except Exception as e:
        print("[!] FFprobe error:", e)

def scan_video_ffprobe(path):
    try:
        cmd = ["ffprobe","-v","quiet","-print_format","json","-show_format","-show_streams", path]
        result = subprocess.run(cmd, capture_output=True, text=True)
        info = json.loads(result.stdout)
        print(json.dumps(info, indent=2))
    except Exception as e:
        print("[!] FFprobe error:", e)

def scan_hachoir(path):
    if createParser is None or extractMetadata is None:
        print("[!] hachoir not installed")
        return
    parser = createParser(path)
    if not parser:
        print("[!] Can't read file!")
        return
    metadata = extractMetadata(parser)
    if metadata:
        print(metadata.exportPlaintext())
    else:
        print("[+] hachoir not finded any metadata")

# -------- MAIN --------

def scan_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in [".pdf"]:
        scan_pdf(path)
    elif ext in [".docx"]:
        scan_docx(path)
    elif ext in [".xlsx"]:
        scan_xlsx(path)
    elif ext in [".pptx"]:
        scan_pptx(path)
    elif ext in [".zip"]:
        scan_zip(path)
    elif ext in [".rar"]:
        scan_rar(path)
    elif ext in [".7z"]:
        scan_7z(path)
    elif ext in [".mp3"]:
        scan_mp3(path)
    elif ext in [".flac"]:
        scan_flac(path)
    elif ext in [".ogg"]:
        scan_ogg(path)
    elif ext in [".m4a"]:
        scan_mp4_audio(path)
    elif ext in [".jpg", ".jpeg", ".tiff", ".tif"]:
        scan_image_exif(path)
    elif ext in [".png", ".webp", ".bmp", ".gif"]:
        scan_image_pillow(path)
    elif ext in [".heic", ".heif"]:
        scan_heic(path)
    elif ext in [".mp4", ".mkv", ".avi", ".mov"]:
        scan_video_ffprobe(path)
    else:
        scan_hachoir(path)

def scan_path(path):
    if os.path.isfile(path):
        scan_file(path)
    elif os.path.isdir(path):
        for root, dirs, files in os.walk(path):
            for file in files:
                print(f"\n--- Scanning: {os.path.join(root,file)} ---")
                scan_file(os.path.join(root, file))
    else:
        print("[!] Path not found")

if __name__ == "__main__":
    path = input("Enter file or folder path: ").strip()
    scan_path(path)